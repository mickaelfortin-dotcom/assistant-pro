"""
Assistant Pro — Serveur Python
Comptes rendus, Gantt, Résumé PDF, Tâches, Gmail
"""
import os, io, json, pickle, base64, hashlib, requests, tempfile
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email import encoders
from flask import Flask, request, jsonify, send_from_directory, session, redirect
from flask_cors import CORS
from dotenv import load_dotenv
from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
from googleapiclient.discovery import build

load_dotenv()
app = Flask(__name__, static_folder=".")
app.secret_key = os.getenv("SECRET_KEY", "supersecret")
app.config['SESSION_COOKIE_SECURE'] = True
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['PERMANENT_SESSION_LIFETIME'] = 30 * 24 * 60 * 60  # 30 jours
CORS(app)

GROQ_API_KEY = os.getenv("GROQ_API_KEY")
GROQ_MODEL   = "llama-3.3-70b-versatile"
GROQ_URL     = "https://api.groq.com/openai/v1/chat/completions"
TOKEN_FILE   = "token.pickle"
APP_PWD_HASH = os.getenv("APP_PASSWORD_HASH", "")

SCOPES = [
    "https://www.googleapis.com/auth/gmail.modify",
    "https://www.googleapis.com/auth/gmail.send",
    "https://www.googleapis.com/auth/calendar",
]

def check_auth(): return session.get("authenticated") == True
def hash_pwd(p): return hashlib.sha256(p.encode()).hexdigest()

def get_creds():
    creds = None
    token_data = os.getenv("GOOGLE_TOKEN")
    if token_data:
        creds = pickle.loads(base64.b64decode(token_data))
    elif os.path.exists(TOKEN_FILE):
        with open(TOKEN_FILE,"rb") as f: creds = pickle.load(f)
    if creds and creds.expired and creds.refresh_token:
        creds.refresh(Request())
    return creds

def gmail_svc(): return build("gmail","v1",credentials=get_creds())
def cal_svc():   return build("calendar","v3",credentials=get_creds())
def google_ok(): c=get_creds(); return c is not None and c.valid

def call_groq(system, messages, max_tokens=2000):
    if not GROQ_API_KEY: raise Exception("Cle Groq manquante")
    msgs = [{"role":"system","content":system}]
    for m in messages: msgs.append({"role":m["role"],"content":m["content"]})
    r = requests.post(GROQ_URL,
        headers={"Authorization":"Bearer "+GROQ_API_KEY,"Content-Type":"application/json"},
        json={"model":GROQ_MODEL,"messages":msgs,"max_tokens":max_tokens,"temperature":0.7},
        timeout=90)
    d = r.json()
    if not r.ok: raise Exception("Erreur Groq: "+d.get("error",{}).get("message",str(r.status_code)))
    return d["choices"][0]["message"]["content"]

# ── LOGIN ──────────────────────────────────────────────────────────
LOGIN = """<!DOCTYPE html><html lang="fr"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1,maximum-scale=1">
<meta name="theme-color" content="#6366f1"><title>Assistant Pro</title>
<style>*{box-sizing:border-box;margin:0;padding:0;}
body{background:#0a0a12;color:#e2e8f0;font-family:'Segoe UI',sans-serif;
display:flex;align-items:center;justify-content:center;min-height:100vh;padding:20px;}
.box{background:#12121e;border:1px solid #1e1e35;border-radius:16px;
padding:40px 32px;width:100%;max-width:380px;text-align:center;}
.icon{font-size:52px;margin-bottom:16px;}
h1{font-size:22px;font-weight:700;margin-bottom:6px;color:#fff;}
p{font-size:13px;color:#6b7280;margin-bottom:28px;}
input{width:100%;background:#0a0a12;border:1px solid #1e1e35;border-radius:10px;
color:#e2e8f0;font-size:16px;padding:13px 16px;outline:none;margin-bottom:14px;
transition:border-color .2s;}input:focus{border-color:#6366f1;}
.btn{width:100%;background:linear-gradient(135deg,#6366f1,#8b5cf6);color:#fff;
border:none;border-radius:10px;font-weight:700;font-size:15px;padding:14px;cursor:pointer;}
.btn:active{opacity:.9;}.err{color:#f87171;font-size:12px;margin-top:8px;}</style>
</head><body><div class="box">
<div class="icon">🤖</div><h1>Assistant Pro</h1>
<p>Entre ton mot de passe pour accéder à ton assistant personnel</p>
<form method="POST" action="/login">
<input type="password" name="password" placeholder="Mot de passe" autofocus>
{error}<button type="submit" class="btn">Connexion →</button></form>
</div></body></html>"""

@app.route("/login", methods=["GET","POST"])
def login():
    if request.method == "POST":
        pwd = request.form.get("password","")
        ok = (APP_PWD_HASH and hash_pwd(pwd)==APP_PWD_HASH) or \
             (not APP_PWD_HASH and pwd==os.getenv("APP_PASSWORD",""))
        if ok:
            session.permanent = True
            session["authenticated"] = True
            return redirect("/")
        return LOGIN.replace("{error}",'<div class="err">❌ Mot de passe incorrect</div>')
    return LOGIN.replace("{error}","")

@app.route("/logout")
def logout(): session.clear(); return redirect("/login")

# ── STATIC ────────────────────────────────────────────────────────
@app.route("/")
def index():
    if not check_auth(): return redirect("/login")
    return send_from_directory(".","index.html")

@app.route("/manifest.json")
def manifest(): return send_from_directory(".","manifest.json")

@app.route("/sw.js")
def sw(): return send_from_directory(".","sw.js")

@app.route("/icons/<path:f>")
def icons(f): return send_from_directory("icons",f)

@app.route("/health")
def health():
    return jsonify({"status":"ok","api_key_configured":bool(GROQ_API_KEY),
                    "google_connected":google_ok(),"authenticated":check_auth()})

# ── IA CHAT ───────────────────────────────────────────────────────
@app.route("/api/chat", methods=["POST"])
def chat():
    if not check_auth(): return jsonify({"error":{"message":"Non autorise"}}),401
    p = request.get_json()
    try:
        text = call_groq(p.get("system","Tu es un assistant expert. Reponds en francais."),
                         p.get("messages",[]), p.get("max_tokens",2000))
        return jsonify({"content":[{"type":"text","text":text}]})
    except Exception as e:
        return jsonify({"error":{"message":str(e)}}),500

# ── UPLOAD & ANALYSE FICHIER ──────────────────────────────────────
@app.route("/api/upload", methods=["POST"])
def upload():
    if not check_auth(): return jsonify({"error":"Non autorise"}),401
    if "file" not in request.files: return jsonify({"error":"Aucun fichier"}),400
    f = request.files["file"]
    fname = f.filename.lower()
    text = ""

    try:
        data = f.read()
        if fname.endswith(".pdf"):
            try:
                import fitz  # PyMuPDF
                doc = fitz.open(stream=data, filetype="pdf")
                for page in doc: text += page.get_text()
                doc.close()
            except ImportError:
                return jsonify({"error":"PyMuPDF non installé"}),500

        elif fname.endswith(".docx"):
            try:
                from docx import Document
                doc = Document(io.BytesIO(data))
                text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            except ImportError:
                return jsonify({"error":"python-docx non installé"}),500

        elif fname.endswith(".xlsx") or fname.endswith(".xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
                for ws in wb.worksheets:
                    text += f"\n--- Feuille: {ws.title} ---\n"
                    for row in ws.iter_rows(values_only=True):
                        row_text = "\t".join([str(c) if c is not None else "" for c in row])
                        if row_text.strip(): text += row_text + "\n"
            except ImportError:
                return jsonify({"error":"openpyxl non installé"}),500

        elif fname.endswith(".txt") or fname.endswith(".md") or fname.endswith(".csv"):
            text = data.decode("utf-8", errors="ignore")
        else:
            return jsonify({"error":"Format non supporté (PDF, DOCX, XLSX, TXT)"}),400

        if not text.strip():
            return jsonify({"error":"Impossible d'extraire le texte du fichier"}),400

        # Tronquer si trop long
        MAX_CHARS = 12000
        truncated = len(text) > MAX_CHARS
        text_sample = text[:MAX_CHARS] + ("\n\n[... contenu tronqué ...]" if truncated else "")

        return jsonify({"text": text_sample, "chars": len(text),
                        "truncated": truncated, "filename": f.filename})

    except Exception as e:
        return jsonify({"error": str(e)}),500

# ── GMAIL ─────────────────────────────────────────────────────────
@app.route("/api/gmail/send", methods=["POST"])
def gmail_send():
    if not check_auth() or not google_ok(): return jsonify({"error":"Non autorise"}),401
    d = request.get_json()
    msg = MIMEMultipart()
    msg["to"] = d["to"]
    msg["subject"] = d["subject"]
    msg.attach(MIMEText(d["body"], "plain", "utf-8"))
    raw = base64.urlsafe_b64encode(msg.as_bytes()).decode()
    gmail_svc().users().messages().send(userId="me", body={"raw":raw}).execute()
    return jsonify({"sent":True,"message":"Email envoyé à "+d["to"]})

@app.route("/api/gmail/stats")
def gmail_stats():
    if not check_auth() or not google_ok(): return jsonify({"error":"Non autorise"}),401
    svc = gmail_svc()
    profile = svc.users().getProfile(userId="me").execute()
    unread = svc.users().messages().list(userId="me",q="is:unread",maxResults=1).execute()
    return jsonify({"email":profile.get("emailAddress",""),
                    "total_messages":profile.get("messagesTotal",0),
                    "total_threads":profile.get("threadsTotal",0),
                    "unread_estimate":unread.get("resultSizeEstimate",0)})

# ── CALENDAR ──────────────────────────────────────────────────────
@app.route("/api/calendar/create", methods=["POST"])
def calendar_create():
    if not check_auth() or not google_ok(): return jsonify({"error":"Non autorise"}),401
    d = request.get_json()
    event = {"summary":d["title"],
             "start":{"dateTime":d["start"],"timeZone":"Europe/Paris"},
             "end":  {"dateTime":d["end"],  "timeZone":"Europe/Paris"}}
    if d.get("description"): event["description"] = d["description"]
    cal_svc().events().insert(calendarId="primary",body=event).execute()
    return jsonify({"created":True,"message":"Événement créé avec succès."})

@app.route("/api/calendar/list")
def calendar_list():
    if not check_auth() or not google_ok(): return jsonify({"error":"Non autorise"}),401
    from datetime import datetime, timezone
    now = datetime.now(timezone.utc).isoformat()
    res = cal_svc().events().list(calendarId="primary",timeMin=now,
          maxResults=10,singleEvents=True,orderBy="startTime").execute()
    events = []
    for e in res.get("items",[]):
        start = e["start"].get("dateTime",e["start"].get("date",""))
        events.append({"id":e["id"],"title":e.get("summary",""),
                       "start":start,"description":e.get("description","")})
    return jsonify({"events":events})

if __name__ == "__main__":
    port = int(os.getenv("PORT",5000))
    print(f"\n{'='*50}\n  Assistant Pro — Serveur\n{'='*50}")
    print(f"  -> http://localhost:{port}")
    print(f"  -> Groq    : {'OK' if GROQ_API_KEY else 'MANQUANTE'}")
    print(f"  -> Google  : {'OK' if google_ok() else 'Non connecté'}")
    print(f"{'='*50}\n")
    app.run(host="0.0.0.0",port=port)
